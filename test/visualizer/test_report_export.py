import io

import pandas as pd

from tricys.visualizer.report.exporter import ReportExporter
from tricys.visualizer.report.templates import TEMPLATES_CATALOG, render_report


def test_templates_catalog_completeness():
    assert len(TEMPLATES_CATALOG) == 3
    template_ids = [t["id"] for t in TEMPLATES_CATALOG]
    assert "simulate" in template_ids
    assert "sweep" in template_ids
    assert "sensitivity" in template_ids

    from tricys.visualizer.report.templates import normalize_template_id

    # Default without template specified -> matches task.type
    assert normalize_template_id(None, "simulate") == "simulate"
    assert normalize_template_id("", "sweep") == "sweep"
    assert normalize_template_id(None, "sensitivity") == "sensitivity"
    # Legacy alias mappings
    assert normalize_template_id("sweep_startup_self_sufficiency") == "sweep"
    assert normalize_template_id("simulate_pulsed_transient") == "simulate"
    assert normalize_template_id("sensitivity_sobol_gsa") == "sensitivity"
    assert normalize_template_id("kpi_metrics") == "sweep"


def test_render_all_templates():
    sample_config = {
        "model": {"modelName": "CFEDR.TritiumCycle"},
        "scenario": {"stopTime": 5000.0, "stepSize": 2.0},
        "task": {
            "type": "sweep",
            "method": "sobol",
            "samples": 64,
            "parameters": {"blanket.TBR": [1.05, 1.25]},
        },
    }
    sample_jobs = [{"id": 1, "blanket.TBR": 1.1}, {"id": 2, "blanket.TBR": 1.2}]
    sample_metrics_df = pd.DataFrame(
        [
            {"job_id": 1, "peak_inventory": 450.2, "settling_time": 1820.0},
            {"job_id": 2, "peak_inventory": 390.5, "settling_time": 1640.0},
        ]
    )

    for t in TEMPLATES_CATALOG:
        t_id = t["id"]
        rendered_md = render_report(
            template_id=t_id,
            config=sample_config,
            jobs=sample_jobs,
            metrics_df=sample_metrics_df,
            ai_insight="AI Recommendation: Maintain TBR above 1.15.",
        )
        assert isinstance(rendered_md, str)
        assert len(rendered_md) > 50
        assert "# TRICYS" in rendered_md


def test_export_markdown():
    text = "# Test Title\n\nSome paragraph text."
    md_output = ReportExporter.export_markdown(text)
    assert md_output.startswith("# Test Title")
    assert md_output.endswith("\n")


def test_export_docx():
    import docx

    md_text = """# TRICYS Executive Report
## 1. Summary
- Bullet item 1 with **bold text**
- Bullet item 2 with `code text`

| Param | Value |
| :--- | :--- |
| TBR | 1.15 |
| Peak | 450.0 |
"""
    docx_bytes = ReportExporter.export_docx(md_text, title="Test Report")
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 500

    # Verify reloading docx
    doc = docx.Document(io.BytesIO(docx_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    assert any("TRICYS Executive Report" in p for p in paragraphs)
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 3


def test_export_pdf():
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz

    md_text = """# TRICYS Engineering PDF Report
## 1. General Info
This is an automated engineering PDF test report.

| Variable | Threshold | Status |
| :--- | :--- | :--- |
| Storage | 500.0 | Normal |

- Safety verification passed
"""
    pdf_bytes = ReportExporter.export_pdf(md_text, title="PDF Test Report")
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 500

    # Verify reloading with fitz
    pdf_doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    assert len(pdf_doc) >= 1
    first_page_text = pdf_doc[0].get_text()
    assert "TRICYS Engineering PDF Report" in first_page_text


def test_export_pdf_with_cjk_characters():
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz

    cjk_md = """# TRICYS 聚变堆全厂氚循环工程仿真报告
## 1. 运行概要
本报告评估聚变堆全流程在役氚滞留量与自持增殖平衡。

| 参数指标 | 设计值 | 状态评估 |
| :--- | :--- | :--- |
| 增殖比 TBR | 1.15 | 合格 |
| 在役储氚 | 450.0 mol | 安全 |

- 质量守恒误差小于 1.0e-5，系统运行平稳。
"""
    pdf_bytes = ReportExporter.export_pdf(cjk_md)
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 500

    pdf_doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    first_page_text = pdf_doc[0].get_text()
    # Ensure Chinese characters are not turned into dots
    assert "聚变堆全厂氚循环工程仿真报告" in first_page_text
    assert "在役储氚" in first_page_text


def test_render_simulate_with_required_tbr_optimization():
    config = {
        "model": {"modelName": "CFEDR.TritiumCycle"},
        "scenario": {"stopTime": 5000.0, "stepSize": 1.0},
        "task": {
            "type": "simulate",
            "metrics": {
                "peak_inventory": {"source": "sds.I[1]", "evaluate": "peak_value"},
                "Required_TBR": {
                    "evaluate": "bisection_search",
                    "parameter": "blanket.TBR",
                    "search_bounds": [1.0, 1.4],
                    "target_metric": {
                        "source": "sds.I[1]",
                        "evaluate": "doubling_time",
                        "threshold": 63072000.0,
                    },
                },
            },
        },
    }
    jobs = [{"job_id": 1, "blanket.TBR": 1.12, "Required_TBR": 1.152}]
    metrics_df = pd.DataFrame(
        [{"job_id": 1, "peak_inventory": 8420.0, "Required_TBR": 1.152}]
    )

    md = render_report(
        template_id="simulate", config=config, jobs=jobs, metrics_df=metrics_df
    )
    assert "任务目标性能与参数寻优指标清单" in md
    assert "Required_TBR" in md
    assert "1.152" in md
